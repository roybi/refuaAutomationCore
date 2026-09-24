#!/usr/bin/env python
"""
elements_to_excel.py - Convert scrape_elements.py's elements.json into an Excel workbook.

Usage:
    python scripts/elements_to_excel.py --json meditik_scrape_output/elements.json --out meditik_scrape_output/elements.xlsx
"""
from __future__ import annotations

import argparse
import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

HEADERS = ["Page", "Element Name", "Type of Element", "Has data-testid", "data-testid Value", "Identification (XPath)"]


def build_workbook(results: dict) -> Workbook:
    wb = Workbook()
    ws = wb.active
    ws.title = "Elements"
    ws.append(HEADERS)

    header_fill = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    header_font = Font(bold=True, color="FFFFFF")
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    row_idx = 2
    for page, data in results.items():
        if data.get("off_domain"):
            continue
        for r in data.get("rows", []):
            ws.append([
                page,
                r.get("name", ""),
                r.get("type", ""),
                r.get("has_testid", ""),
                r.get("testid", ""),
                r.get("xpath", ""),
            ])
            row_idx += 1

    # Right-to-left friendly + readable column widths
    ws.sheet_view.rightToLeft = True
    widths = [22, 45, 18, 14, 40, 55]
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{row_idx - 1}"

    return wb


def main(argv=None):
    ap = argparse.ArgumentParser(description="Convert elements.json to an Excel workbook.")
    ap.add_argument("--json", default="meditik_scrape_output/elements.json", help="Path to elements.json")
    ap.add_argument("--out", default="meditik_scrape_output/elements.xlsx", help="Output .xlsx path")
    args = ap.parse_args(argv)

    json_path = Path(args.json)
    results = json.loads(json_path.read_text(encoding="utf-8"))

    wb = build_workbook(results)
    out_path = Path(args.out)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out_path))
    print(f"XLSX {out_path}")


if __name__ == "__main__":
    raise SystemExit(main())
